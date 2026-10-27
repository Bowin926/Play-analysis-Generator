# -*- coding: utf-8 -*-
"""生成剧本分析 Word + 场景/角色提示词 txt 的辅助脚本。

用法示例：
    python generate_docs.py --play 哈姆雷特 --image-style 手绘
"""
import argparse
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn


def styled_run(paragraph, text, bold=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return run


def h(doc, level, text):
    p = doc.add_heading(level=level)
    styled_run(p, text, bold=True)
    p.paragraph_format.space_after = Pt(6)


def para(doc, text):
    p = doc.add_paragraph()
    styled_run(p, text)


def bullet(doc, lead, text=""):
    p = doc.add_paragraph(style="List Bullet")
    if lead:
        styled_run(p, lead + "：", bold=True)
    if text:
        styled_run(p, text)


def write_docx(play_name, analysis):
    doc = Document()
    h(doc, 0, f"《{play_name}》剧本分析")

    h(doc, 1, "一、形象种子")
    bullet(doc, "形象种子", analysis["seed"])
    para(doc, "一句话解释：" + analysis["seed_explanation"])

    h(doc, 1, "二、故事梗概")
    h(doc, 2, "整体概括")
    para(doc, analysis["summary"])
    h(doc, 2, "分幕描述")
    for act, text in analysis["acts"].items():
        h(doc, 2, act)
        para(doc, text)

    h(doc, 1, "三、场景描述")
    for act, text in analysis["scenes"].items():
        h(doc, 2, act)
        para(doc, text)

    h(doc, 1, "四、矛盾冲突")
    h(doc, 2, "全剧核心矛盾")
    para(doc, analysis["core_conflict"])
    for act, (major, minor) in analysis["conflicts"].items():
        p = doc.add_paragraph()
        styled_run(p, act, bold=True)
        bullet(doc, "主要矛盾", major)
        bullet(doc, "次要矛盾", minor)

    h(doc, 1, "五、角色分析")
    for char in analysis["characters"]:
        h(doc, 2, "角色：" + char["name"])
        bullet(doc, "角色的形象种子", char["seed"])
        bullet(doc, "角色形象描述", char["description"])
        bullet(doc, "角色的性格", char["character"])
        p = doc.add_paragraph()
        styled_run(p, "角色的核心矛盾", bold=True)
        bullet(doc, "内在矛盾", char["inner_conflict"])
        bullet(doc, "外在矛盾", char["outer_conflict"])
        bullet(doc, "角色的结局", char["ending"])
        p = doc.add_paragraph()
        styled_run(p, "角色关系", bold=True)
        bullet(doc, "主要关系", char["relationships"])

    out = f"{play_name}剧本分析.docx"
    doc.save(out)
    print("saved", out)


def write_scene_prompts(play_name, scenes, image_ratio="16:9", venue="镜框式舞台"):
    out = f"{play_name}场景画面提示词.txt"
    lines = [f"《{play_name}》场景画面提示词", ""]
    for act, text in scenes.items():
        lines.append(act + "：")
        lines.append("剧场舞台概念图，机位要求正视角构图、正面平视、居中呈现，" +
                     f"画面比例 {image_ratio}，演出场地为{venue}。{text}")
        lines.append("")
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print("saved", out)


def write_character_prompts(play_name, characters, style=None):
    out = f"{play_name}角色形象提示词.txt"
    lines = [f"《{play_name}》角色形象提示词",
             "通用规范：白底角色效果图，全身像，角色独立居中，画面比例 3:4，背景统一为纯白，不出现舞台、道具或场景说明文字。", ""]

    def base(selected_style):
        if selected_style == "写实":
            return ("白底角色效果图，写实风格，电影级打光，材质真实，皮肤、布料和金属质感明确；"
                    "角色全身独立居中，全身像，画面比例 3:4，背景统一为纯白。")
        return ("白底角色效果图，手绘风格，铅笔与水彩质感，线条清晰但不僵硬，色彩干净；"
                "角色全身独立居中，全身像，画面比例 3:4，背景统一为纯白。")

    styles = [style] if style else ["手绘", "写实"]
    for current in styles:
        lines.append(f"==== {current}风格版 ====")
        lines.append("")
        for char in characters:
            lines.append(char["name"] + "：")
            lines.append(base(current) + "角色形象与服装：" + char["costume"])
            lines.append("")
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print("saved", out)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--play", required=True, help="剧本名，例如 哈姆雷特")
    parser.add_argument("--image-style", default=None, help="角色图风格：手绘 或 写实")
    parser.add_argument("--scene-ratio", default="16:9")
    parser.add_argument("--venue", default="镜框式舞台")
    args = parser.parse_args()

    # 由代理在调用前填充实际内容。示例留空，避免生成错误文档。
    analysis = {
        "seed": "",
        "seed_explanation": "",
        "summary": "",
        "acts": {},
        "scenes": {},
        "core_conflict": "",
        "conflicts": {},
        "characters": [],
    }
    write_docx(args.play, analysis)
    write_scene_prompts(args.play, analysis["scenes"], args.scene_ratio, args.venue)
    write_character_prompts(args.play, analysis["characters"], args.image_style)


if __name__ == "__main__":
    main()
